#!/usr/bin/env python3
"""Turn client answers typed into the ledger spreadsheet into client_confirmed claims for version 2.

Usage:
  python scripts/import_answers.py --ledger <client>-claims-ledger-INTERNAL.xlsx --data report-data.json --out report-data-v2.json

Reads the "Answer from client" column on the Questions sheet. Each non-empty answer becomes a new claim
(evidence client_confirmed, status included, section from the question's topic) and the question is marked
answered. You still write the version 2 sections and blocks that cite these claims; this script only adds the
claims so nothing confirmed in the meeting is lost or paraphrased loosely.
"""
import argparse, json
from openpyxl import load_workbook

TOPIC_SECTION = {"Goals and priorities": "goals", "Customers and decision makers": "profiles", "Offerings and pricing": "offerings",
                 "Competitors they lose to": "competitors", "Brand and voice": "positioning", "Marketing and measurement": "plan",
                 "Constraints and topics to avoid": "plan"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ledger", required=True); ap.add_argument("--data", required=True); ap.add_argument("--out", required=True)
    ap.add_argument("--meeting-date", default="")
    a = ap.parse_args()
    data = json.load(open(a.data, encoding="utf-8"))
    ws = load_workbook(a.ledger, data_only=True)["Questions"]
    headers = [str(c.value) for c in ws[1]]
    qi, ai = headers.index("ID"), headers.index("Answer from client")
    ti = headers.index("Topic"); qq = headers.index("Question")
    existing = {c["id"] for c in data["claims"]}
    n = max([int(c["id"][1:]) for c in data["claims"] if c["id"][1:].isdigit()] + [0])
    added = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        qid, ans = row[qi], row[ai]
        if not ans or not str(ans).strip():
            continue
        n += 1
        cid = f"C{n:03d}"
        claim = {"id": cid, "section": TOPIC_SECTION.get(str(row[ti]), "goals"),
                 "text": f"The client confirmed: {str(ans).strip()}", "quote": "",
                 "source_url": "", "source_title": f"Client meeting{(' ' + a.meeting_date) if a.meeting_date else ''}",
                 "evidence": "client_confirmed", "status": "included", "note": f"Answer to {qid}: {row[qq]}"}
        data["claims"].append(claim); added.append(claim)
        for q in data.get("questions", []):
            if q["id"] == qid:
                q["answered"] = True; q["answer"] = str(ans).strip()
    data["version"] = "v2"
    json.dump(data, open(a.out, "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    print(json.dumps({"added": added, "remaining_questions": [q["id"] for q in data.get("questions", []) if not q.get("answered")]}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
